"""
OCR Service
Convert scanned PDFs/images to searchable PDFs or extract text
"""

from pathlib import Path
from typing import List, Tuple, Literal
import logging
import io
import os
import shutil

import pytesseract
from PIL import Image
from pdf2image import convert_from_path
from pypdf import PdfReader
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from PyPDF2 import PdfWriter, PdfReader as PyPDF2Reader

# Configure Tesseract path for Windows - try multiple locations
def find_tesseract():
    """Find Tesseract executable"""
    # Try common Windows paths
    possible_paths = [
        r'C:\Program Files\Tesseract-OCR\tesseract.exe',
        r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
        r'C:\Tesseract-OCR\tesseract.exe',
    ]
    
    for path in possible_paths:
        if os.path.exists(path):
            return path
    
    # Try to find in PATH
    tesseract_path = shutil.which('tesseract')
    if tesseract_path:
        return tesseract_path
    
    return None

# Set Tesseract path
tesseract_path = find_tesseract()
if tesseract_path:
    pytesseract.pytesseract.tesseract_cmd = tesseract_path
else:
    # Set default path - will fail with helpful error message later
    pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

logger = logging.getLogger(__name__)

OutputFormat = Literal["searchable_pdf", "text", "docx"]


class OcrService:
    """Service for OCR operations on PDFs and images"""
    
    @staticmethod
    def verify_tesseract():
        """Verify Tesseract is installed"""
        tesseract_cmd = pytesseract.pytesseract.tesseract_cmd
        if not os.path.exists(tesseract_cmd):
            raise Exception(
                f"Tesseract-OCR is not installed at {tesseract_cmd}.\n"
                "Please download and install from:\n"
                "https://digi.bib.uni-mannheim.de/tesseract/tesseract-ocr-w64-setup-5.3.3.20231005.exe\n"
                "Install to the default path: C:\\Program Files\\Tesseract-OCR"
            )
    
    @staticmethod
    def detect_pdf_has_text(pdf_path: Path, sample_pages: int = 3) -> bool:
        """
        Detect if PDF already contains searchable text
        
        Args:
            pdf_path: Path to PDF file
            sample_pages: Number of pages to sample (default: 3)
            
        Returns:
            True if PDF has text, False if scanned/image-only
        """
        try:
            reader = PdfReader(pdf_path)
            total_pages = len(reader.pages)
            pages_to_check = min(sample_pages, total_pages)
            
            text_found = False
            for i in range(pages_to_check):
                text = reader.pages[i].extract_text().strip()
                if text and len(text) > 50:  # Meaningful text threshold
                    text_found = True
                    break
            
            logger.info(f"PDF text detection: {'Has text' if text_found else 'Scanned/No text'}")
            return text_found
            
        except Exception as e:
            logger.error(f"Error detecting PDF text: {e}")
            return False
    
    @staticmethod
    def pdf_to_images(pdf_path: Path, dpi: int = 300) -> List[Image.Image]:
        """
        Convert PDF pages to PIL images
        
        Args:
            pdf_path: Path to PDF file
            dpi: Resolution for conversion (default: 300)
            
        Returns:
            List of PIL Image objects
        """
        try:
            logger.info(f"Converting PDF to images at {dpi} DPI")
            # Specify Poppler path for Windows
            poppler_path = r'C:\Program Files\poppler\Library\bin'
            images = convert_from_path(str(pdf_path), dpi=dpi, poppler_path=poppler_path)
            logger.info(f"Converted {len(images)} pages to images")
            return images
            
        except Exception as e:
            logger.error(f"Error converting PDF to images: {e}")
            raise Exception(f"Failed to convert PDF to images: {str(e)}")
    
    @staticmethod
    def extract_text_from_image(image: Image.Image, lang: str = 'eng') -> str:
        """
        Extract text from image using Tesseract OCR
        
        Args:
            image: PIL Image object
            lang: Language code for OCR (default: 'eng')
            
        Returns:
            Extracted text
        """
        try:
            # Configure Tesseract for better accuracy
            custom_config = r'--oem 3 --psm 3'  # LSTM OCR, automatic page segmentation
            text = pytesseract.image_to_string(image, lang=lang, config=custom_config)
            return text.strip()
            
        except Exception as e:
            logger.error(f"Error extracting text from image: {e}")
            return ""
    
    @staticmethod
    def create_searchable_pdf(
        images: List[Image.Image],
        output_path: Path,
        lang: str = 'eng'
    ) -> Path:
        """
        Create searchable PDF from images with OCR text layer
        
        Args:
            images: List of PIL Image objects
            output_path: Path for output PDF
            lang: Language code for OCR
            
        Returns:
            Path to searchable PDF
        """
        try:
            logger.info(f"Creating searchable PDF from {len(images)} images")
            
            pdf_writer = PdfWriter()
            
            for i, image in enumerate(images, 1):
                logger.info(f"Processing page {i}/{len(images)}")
                
                # Create temporary file for OCR PDF
                temp_pdf = output_path.parent / f"temp_ocr_page_{i}.pdf"
                
                # Run OCR and create PDF with text layer
                pdf_bytes = pytesseract.image_to_pdf_or_hocr(
                    image,
                    lang=lang,
                    extension='pdf',
                    config='--oem 3 --psm 3'
                )
                
                # Write temporary PDF
                with open(temp_pdf, 'wb') as f:
                    f.write(pdf_bytes)
                
                # Read and add to final PDF
                with open(temp_pdf, 'rb') as f:
                    temp_reader = PyPDF2Reader(f)
                    pdf_writer.add_page(temp_reader.pages[0])
                
                # Clean up temp file
                try:
                    temp_pdf.unlink()
                except:
                    pass
            
            # Write final PDF
            with open(output_path, 'wb') as f:
                pdf_writer.write(f)
            
            logger.info(f"Successfully created searchable PDF: {output_path.name}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error creating searchable PDF: {e}")
            raise Exception(f"Failed to create searchable PDF: {str(e)}")
    
    @staticmethod
    def extract_text_to_file(
        images: List[Image.Image],
        output_path: Path,
        format: Literal["text", "docx"],
        lang: str = 'eng'
    ) -> Path:
        """
        Extract text from images and save as text or docx
        
        Args:
            images: List of PIL Image objects
            output_path: Path for output file
            format: Output format ('text' or 'docx')
            lang: Language code for OCR
            
        Returns:
            Path to output file
        """
        try:
            logger.info(f"Extracting text to {format} from {len(images)} images")
            
            all_text = []
            for i, image in enumerate(images, 1):
                logger.info(f"OCR processing page {i}/{len(images)}")
                text = OcrService.extract_text_from_image(image, lang)
                all_text.append(text)
            
            if format == "text":
                # Save as plain text
                with open(output_path, 'w', encoding='utf-8') as f:
                    for i, text in enumerate(all_text, 1):
                        if len(all_text) > 1:
                            f.write(f"\n\n--- Page {i} ---\n\n")
                        f.write(text)
                        f.write("\n")
                
                logger.info(f"Successfully created text file: {output_path.name}")
                
            elif format == "docx":
                # Save as Word document
                doc = Document()
                doc.add_heading('OCR Text Extract', 0)
                
                for i, text in enumerate(all_text, 1):
                    if len(all_text) > 1:
                        doc.add_heading(f'Page {i}', level=1)
                    
                    # Add text with paragraphs
                    for paragraph in text.split('\n'):
                        if paragraph.strip():
                            doc.add_paragraph(paragraph)
                
                doc.save(output_path)
                logger.info(f"Successfully created docx file: {output_path.name}")
            
            return output_path
            
        except Exception as e:
            logger.error(f"Error extracting text to {format}: {e}")
            raise Exception(f"Failed to extract text to {format}: {str(e)}")
    
    @staticmethod
    def process_ocr(
        input_file: Path,
        output_file: Path,
        output_format: OutputFormat,
        force_ocr: bool = False,
        lang: str = 'eng'
    ) -> Tuple[Path, bool]:
        """
        Main OCR processing function
        
        Args:
            input_file: Path to input PDF or image
            output_file: Path for output file
            output_format: Output format ('searchable_pdf', 'text', 'docx')
            force_ocr: Force OCR even if PDF has text
            lang: Language code for OCR
            
        Returns:
            Tuple of (output_path, was_ocr_performed)
        """
        try:
            # Verify Tesseract is installed first
            OcrService.verify_tesseract()
            
            is_pdf = input_file.suffix.lower() == '.pdf'
            needs_ocr = force_ocr
            
            # Check if PDF needs OCR
            if is_pdf and not force_ocr:
                has_text = OcrService.detect_pdf_has_text(input_file)
                needs_ocr = not has_text
                
                if not needs_ocr and output_format == "searchable_pdf":
                    # PDF already has text, just copy it
                    logger.info("PDF already has text, no OCR needed")
                    import shutil
                    shutil.copy(input_file, output_file)
                    return output_file, False
            
            # Convert to images
            if is_pdf:
                images = OcrService.pdf_to_images(input_file)
            else:
                # Single image file
                images = [Image.open(input_file)]
            
            # Process based on output format
            if output_format == "searchable_pdf":
                result = OcrService.create_searchable_pdf(images, output_file, lang)
            else:
                result = OcrService.extract_text_to_file(
                    images, output_file, output_format, lang
                )
            
            return result, True
            
        except Exception as e:
            logger.error(f"Error processing OCR: {e}")
            raise Exception(f"OCR processing failed: {str(e)}")

# okay
